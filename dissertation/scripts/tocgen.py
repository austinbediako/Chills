#!/usr/bin/env python3
"""Pass-2: populate static TOC/LOF/LOT in the docx using page numbers
measured from the pass-1 PDF render.
Usage: python3 tocgen.py KBlog-Dissertation-formatted.docx pass1.pdf out.docx
"""
import re, sys
from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
sys.path.insert(0, '/Users/kaeytee/Library/Python/3.9/lib/python/site-packages')
from pypdf import PdfReader

DOCX, PDF, OUT = sys.argv[1], sys.argv[2], sys.argv[3]

doc = Document(DOCX)
pdf = PdfReader(PDF)
pages = [' '.join((p.extract_text() or '').split()) for p in pdf.pages]

# locate body start via body text unique to CH1 (not the TOC listing)
body_start = next(i for i, t in enumerate(pages) if 'This dissertation reports' in t)

FRONT_HEADS = ('DECLARATION', 'ABSTRACT', 'DEDICATION', 'ACKNOWLEDGEMENT',
               'TABLE OF CONTENTS', 'LIST OF FIGURES', 'LIST OF TABLES', 'LIST OF ABBREVIATIONS')

def roman(n):
    vals = [(10,'x'),(9,'ix'),(5,'v'),(4,'iv'),(1,'i')]
    s = ''
    for v, sym in vals:
        while n >= v: s += sym; n -= v
    return s

def disp(i):
    return roman(i + 1) if i < body_start else str(i - body_start + 1)

def find_page(snippet):
    sn = ' '.join(snippet.split())[:60]
    front = any(sn.startswith(h) for h in FRONT_HEADS)
    rng = range(0, body_start) if front else range(body_start, len(pages))
    if front:
        # list-headings may sit mid-page after the previous list; they are
        # identified by their first entry following within ~250 chars
        followers = {'LIST OF FIGURES': 'Figure 3.1', 'LIST OF TABLES': 'Table 2.1',
                     'LIST OF ABBREVIATIONS': 'Abbreviation'}
        for h, f in followers.items():
            if sn.startswith(h):
                for i in rng:
                    pos = pages[i].find(sn[:40])
                    while pos >= 0:
                        if pages[i].find(f, pos, pos + 250) >= 0:
                            return i
                        pos = pages[i].find(sn[:40], pos + 1)
                return None
        # other front headings appear at page start (not inside TOC listing)
        for i in rng:
            pos = pages[i].find(sn[:40])
            if 0 <= pos <= 80:
                return i
        return None
    for i in rng:
        if sn[:40] in pages[i]:
            return i
    for i in rng:
        if sn[:30] in pages[i]:
            return i
    return None

# collect entries from docx
toc, lof, lot = [], [], []
for p in doc.paragraphs:
    st = p.style.name
    txt = p.text.strip()
    if not txt:
        continue
    if st == 'Heading 1':
        toc.append((1, txt.replace('\n', ' ')))
    elif st == 'Heading 2':
        toc.append((2, txt))
    elif st == 'Heading 3':
        toc.append((3, txt))
    elif st == 'Image Caption':
        lof.append(txt)
    elif st == 'Table Caption':
        lot.append(txt)

def entry_paragraph(text, page):
    return (text, disp(page) if page is not None else '')

def write_entries(field_par, entries, doc):
    """Replace field-paragraph siblings: insert static entries after the field para.
    First removes any previously-inserted static entries (paragraphs between the
    field para and the next Heading 1) so the pass is idempotent."""
    from docx.text.paragraph import Paragraph
    from docx.oxml import OxmlElement
    # cleanup: delete paragraphs after field para until next Heading 1
    nxt = field_par._p.getnext()
    while nxt is not None:
        pstyle = nxt.find('.//' + qn('w:pStyle'))
        if pstyle is not None and pstyle.get(qn('w:val')) == 'Heading1':
            break
        # also stop at paragraphs containing a TOC instrText (next field)
        if nxt.findall('.//' + qn('w:instrText')):
            break
        nxt2 = nxt.getnext()
        nxt.getparent().remove(nxt)
        nxt = nxt2
    anchor = field_par
    for level, text, page in entries:
        np = OxmlElement('w:p')
        anchor._p.addnext(np)
        pel = Paragraph(np, field_par._parent)
        run = pel.add_run(text)
        run.font.size = Pt(12); run.font.name = 'Times New Roman'
        pel.paragraph_format.left_indent = Inches(0.25 * (level - 1))
        pel.paragraph_format.line_spacing = 1.5
        pel.paragraph_format.tab_stops.add_tab_stop(Inches(6.3), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        pel.add_run('\t')
        run2 = pel.add_run(page)
        run2.font.size = Pt(12); run2.font.name = 'Times New Roman'
        anchor = pel

# find the three field paragraphs (contain instrText)
def field_instr(p):
    return ' '.join(el.text or '' for el in p._p.findall('.//' + qn('w:instrText')))

for p in doc.paragraphs:
    instr = field_instr(p)
    if not instr:
        continue
    if 'TOC \\o' in instr:
        entries = [(lv, t, entry_paragraph(t, find_page(t))[1]) for lv, t in toc]
        write_entries(p, entries, doc)
        # blank out placeholder run text
        for r in p.runs:
            if 'Update Field' in r.text: r.text = ''
    elif 'Image Caption' in instr:
        entries = [(1, t, entry_paragraph(t, find_page(t))[1]) for t in lof]
        write_entries(p, entries, doc)
        for r in p.runs:
            if 'Update Field' in r.text: r.text = ''
    elif 'Table Caption' in instr:
        entries = [(1, t, entry_paragraph(t, find_page(t))[1]) for t in lot]
        write_entries(p, entries, doc)
        for r in p.runs:
            if 'Update Field' in r.text: r.text = ''

doc.save(OUT)
print('wrote', OUT)
